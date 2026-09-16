#!/usr/bin/env python3
"""
SAFNEX NOVA PowerPoint Generator
Creates professional presentation from content
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 1920px
    prs.slide_height = Inches(7.5)    # 1080px

    # Define colors
    PRIMARY = RGBColor(102, 126, 234)  # #667eea
    SECONDARY = RGBColor(118, 75, 162)  # #764ba2
    ACCENT = RGBColor(85, 207, 255)     # #55cfff
    SUCCESS = RGBColor(34, 197, 94)     # #22c55e
    WARNING = RGBColor(251, 146, 60)    # #fb923c

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    title = slide.shapes.title
    left = Inches(2)
    top = Inches(2.5)
    width = Inches(9.333)
    height = Inches(2)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = "🛡️ SAFNEX NOVA"
    p = tf.paragraphs[0]
    p.font.size = Pt(66)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    left = Inches(2)
    top = Inches(4.5)
    txBox = slide.shapes.add_textbox(left, top, width, Inches(1))
    tf = txBox.text_frame
    tf.text = "AI-Powered Unified Security Analysis Platform"
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

    # Date
    left = Inches(2)
    top = Inches(6)
    txBox = slide.shapes.add_textbox(left, top, width, Inches(0.5))
    tf = txBox.text_frame
    tf.text = "Team Presentation | September 16, 2026"
    p = tf.paragraphs[0]
    p.font.size = Pt(18)
    p.alignment = PP_ALIGN.CENTER

    # Slide 2: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1))
    tf = txBox.text_frame
    tf.text = "The Problem"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    # Stats
    stats = [
        ("85%", "Cyberattacks start with phishing"),
        ("$10B+", "Annual SMS scam losses"),
        ("587%", "QR code phishing increase"),
        ("∞", "Deepfakes now accessible to criminals")
    ]

    row = 0
    for i, (number, desc) in enumerate(stats):
        col = i % 2
        left = Inches(1.5 + col * 5.5)
        top = Inches(2 + row * 2)

        # Number
        txBox = slide.shapes.add_textbox(left, top, Inches(4), Inches(0.8))
        tf = txBox.text_frame
        tf.text = number
        p = tf.paragraphs[0]
        p.font.size = Pt(60)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.alignment = PP_ALIGN.CENTER

        # Description
        txBox = slide.shapes.add_textbox(left, top + Inches(1), Inches(4), Inches(0.5))
        tf = txBox.text_frame
        tf.text = desc
        p = tf.paragraphs[0]
        p.font.size = Pt(16)
        p.alignment = PP_ALIGN.CENTER

        if col == 1:
            row += 1

    # Bottom text
    txBox = slide.shapes.add_textbox(Inches(2), Inches(6.5), Inches(9.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "Modern threats are multi-modal. We need a unified solution."
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.color.rgb = WARNING
    p.alignment = PP_ALIGN.CENTER

    # Slide 3: 7 Tools Overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "SAFNEX NOVA"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(11.333), Inches(0.5))
    tf = txBox.text_frame
    tf.text = "7 Security Tools in 1 Platform"
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

    # Tools grid
    tools = [
        ("🔗", "Link Intelligence", "Phishing detection"),
        ("📞", "Phone Validator", "Fraud detection"),
        ("💬", "Message Detector", "Scam analysis"),
        ("📱", "QR Scanner", "Safe QR codes"),
        ("🖼️", "Image AI", "Deepfake images"),
        ("🎤", "Voice AI", "Audio deepfakes"),
        ("🎥", "Video AI", "Video deepfakes")
    ]

    for i, (icon, name, desc) in enumerate(tools):
        row = i // 3
        col = i % 3
        left = Inches(0.8 + col * 4.1)
        top = Inches(2.5 + row * 1.8)

        # Icon
        txBox = slide.shapes.add_textbox(left, top, Inches(3.5), Inches(0.5))
        tf = txBox.text_frame
        tf.text = icon
        p = tf.paragraphs[0]
        p.font.size = Pt(40)
        p.alignment = PP_ALIGN.CENTER

        # Name
        txBox = slide.shapes.add_textbox(left, top + Inches(0.6), Inches(3.5), Inches(0.4))
        tf = txBox.text_frame
        tf.text = name
        p = tf.paragraphs[0]
        p.font.size = Pt(18)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

        # Description
        txBox = slide.shapes.add_textbox(left, top + Inches(1), Inches(3.5), Inches(0.3))
        tf = txBox.text_frame
        tf.text = desc
        p = tf.paragraphs[0]
        p.font.size = Pt(14)
        p.alignment = PP_ALIGN.CENTER

    # Slide 4: Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "System Architecture"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    # Architecture diagram
    arch_text = """
┌─────────────────────────────────┐
│    WEB APP + CHROME EXTENSION   │
└───────────────┬─────────────────┘
                │ HTTPS
┌───────────────▼─────────────────┐
│   NODE.JS + EXPRESS BACKEND     │
│        (4,167 lines)            │
└───────────────┬─────────────────┘
                │
    ┌───────────┴────────┬────────────┐
    │                    │            │
┌───▼────────┐   ┌──────▼────┐  ┌───▼──────┐
│  OpenAI    │   │  Twilio   │  │NumVerify │
│  GPT-4o    │   │  Lookup   │  │   API    │
└────────────┘   └───────────┘  └──────────┘
    """

    txBox = slide.shapes.add_textbox(Inches(3), Inches(2), Inches(7.333), Inches(3))
    tf = txBox.text_frame
    tf.text = arch_text
    p = tf.paragraphs[0]
    p.font.name = 'Courier New'
    p.font.size = Pt(14)

    # Tech stack
    tech = ["Node.js", "Express", "OpenAI GPT-4o", "Tesseract.js", "Twilio", "Manifest V3"]
    tech_text = "  •  ".join(tech)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(11.333), Inches(1))
    tf = txBox.text_frame
    tf.text = tech_text
    p = tf.paragraphs[0]
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.CENTER

    # Slide 5: Link Intelligence
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "🔗 Link Intelligence Analyzer"
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(11.333), Inches(0.5))
    tf = txBox.text_frame
    tf.text = "How It Works"
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

    steps = [
        "1. URL Normalization - Convert to standard format",
        "2. SSRF Protection - Block private IPs (10.x, 192.168.x, localhost)",
        "3. DNS Resolution - Verify domain exists",
        "4. Known Service Detection - Whitelist trusted services",
        "5. AI Analysis - GPT-4o deep inspection",
        "6. Risk Score - Calculate 0-100% threat level"
    ]

    txBox = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9.333), Inches(4))
    tf = txBox.text_frame

    for step in steps:
        p = tf.add_paragraph()
        p.text = step
        p.font.size = Pt(20)
        p.level = 0
        p.space_before = Pt(12)

    # Slide 6: Link Scoring Formula
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "Link Risk Score Formula"
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    formula = """Base Score = 0

IF (private IP)      → +100 (CRITICAL)
IF (suspicious TLD)  → +40
IF (URL shortener)   → +35
IF (misspelled brand)→ +45
IF (IP in URL)       → +50
IF (no HTTPS)        → +25
IF (known safe)      → -30

AI Override: GPT-4o adjusts with reasoning

Final = CLAMP(Score, 0, 100)"""

    txBox = slide.shapes.add_textbox(Inches(2.5), Inches(2), Inches(8.333), Inches(3.5))
    tf = txBox.text_frame
    tf.text = formula
    p = tf.paragraphs[0]
    p.font.name = 'Courier New'
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT

    # Stats
    txBox = slide.shapes.add_textbox(Inches(3), Inches(6), Inches(7.333), Inches(1))
    tf = txBox.text_frame
    tf.text = "Accuracy: 92-95%  •  Speed: 2-4 seconds"
    p = tf.paragraphs[0]
    p.font.size = Pt(22)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Slide 7: Detection Accuracy
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(0.8))
    tf = txBox.text_frame
    tf.text = "📊 Detection Accuracy"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    accuracies = [
        ("🔗 Link Intelligence", "92-95%"),
        ("📞 Phone Validator", "90-98%"),
        ("🖼️ Image AI", "85-95%"),
        ("🎤 Voice AI", "82-90%"),
        ("🎥 Video AI", "80-88%"),
        ("💬 Message Scam", "88-94%")
    ]

    for i, (tool, accuracy) in enumerate(accuracies):
        row = i // 2
        col = i % 2
        left = Inches(1 + col * 6)
        top = Inches(2 + row * 1.5)

        # Tool name
        txBox = slide.shapes.add_textbox(left, top, Inches(5), Inches(0.4))
        tf = txBox.text_frame
        tf.text = tool
        p = tf.paragraphs[0]
        p.font.size = Pt(18)
        p.font.bold = True

        # Accuracy
        txBox = slide.shapes.add_textbox(left, top + Inches(0.5), Inches(5), Inches(0.6))
        tf = txBox.text_frame
        tf.text = accuracy
        p = tf.paragraphs[0]
        p.font.size = Pt(48)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.alignment = PP_ALIGN.CENTER

    # Slide 8: Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(9.333), Inches(1))
    tf = txBox.text_frame
    tf.text = "🛡️ Thank You"
    p = tf.paragraphs[0]
    p.font.size = Pt(66)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.CENTER

    txBox = slide.shapes.add_textbox(Inches(2), Inches(3.5), Inches(9.333), Inches(1))
    tf = txBox.text_frame
    tf.text = "SAFNEX NOVA"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

    txBox = slide.shapes.add_textbox(Inches(2), Inches(4.8), Inches(9.333), Inches(1))
    tf = txBox.text_frame
    tf.text = "Protecting the Digital World, One Analysis at a Time"
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER

    txBox = slide.shapes.add_textbox(Inches(2), Inches(6.2), Inches(9.333), Inches(0.5))
    tf = txBox.text_frame
    tf.text = "safnex-nova.onrender.com"
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

    # Save
    output_path = r"C:\Users\MUTHU KRISHNAN\Downloads\SAFNEX-NOVA-Presentation.pptx"
    prs.save(output_path)
    print(f"✓ PowerPoint created: {output_path}")
    return output_path

if __name__ == "__main__":
    try:
        path = create_presentation()
        print(f"\n✓ SUCCESS!")
        print(f"  File: {path}")
        print(f"  Location: Downloads folder")
        print(f"  Size: {len(open(path, 'rb').read()) / 1024:.1f} KB")
    except Exception as e:
        print(f"✗ ERROR: {e}")
        import traceback
        traceback.print_exc()
